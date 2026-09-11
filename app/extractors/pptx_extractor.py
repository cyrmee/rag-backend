import subprocess
import tempfile
import uuid
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.extractors.types import ExtractedImage, ExtractionResult, TextChunk


def _chart_to_markdown(chart) -> str | None:
    """Prefer the chart's underlying data table over a visual screenshot -
    exact numbers are cheaper and more precise than a vision model reading
    bar heights (Phase 2 judgment call)."""
    try:
        plot = chart.plots[0]
        categories = [str(c) for c in plot.categories]
        series = list(plot.series)
        if not categories or not series:
            return None

        header = "| Category | " + " | ".join(s.name or f"Series {i+1}" for i, s in enumerate(series)) + " |"
        sep = "| --- | " + " | ".join("---" for _ in series) + " |"
        rows = []
        for i, cat in enumerate(categories):
            values = [str(s.values[i]) if i < len(s.values) else "" for s in series]
            rows.append(f"| {cat} | " + " | ".join(values) + " |")
        return "\n".join([header, sep] + rows)
    except Exception:
        return None


def _render_slide_images(file_path: str, slide_count: int) -> dict[int, bytes]:
    """Fallback for native charts with no readable data table: render each
    slide to an image via LibreOffice headless so the vision model sees the
    chart as it visually appears."""
    with tempfile.TemporaryDirectory() as tmp:
        pdf_path = Path(tmp) / "slides.pdf"
        subprocess.run(
            [
                "soffice", "--headless", "--convert-to", "pdf",
                "--outdir", tmp, file_path,
            ],
            check=True, capture_output=True, timeout=120,
        )
        converted = next(Path(tmp).glob("*.pdf"))

        import pymupdf as fitz
        doc = fitz.open(str(converted))
        try:
            images = {}
            for i, page in enumerate(doc):
                pix = page.get_pixmap(dpi=150)
                images[i] = pix.tobytes("png")
            return images
        finally:
            doc.close()


def extract_pptx(file_path: str) -> ExtractionResult:
    prs = Presentation(file_path)
    text_chunks: list[TextChunk] = []
    images: list[ExtractedImage] = []
    slides_needing_render: set[int] = set()

    for slide_index, slide in enumerate(prs.slides):
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                text_chunks.append(TextChunk(content=shape.text_frame.text, page_number=slide_index + 1))

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                images.append(
                    ExtractedImage(
                        image_id=str(uuid.uuid4()),
                        image_bytes=shape.image.blob,
                        page_number=slide_index + 1,
                        bbox=None,
                    )
                )
            elif shape.has_chart:
                markdown = _chart_to_markdown(shape.chart)
                if markdown:
                    text_chunks.append(
                        TextChunk(content=markdown, source_type="chart_data", page_number=slide_index + 1)
                    )
                else:
                    slides_needing_render.add(slide_index)

        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                text_chunks.append(
                    TextChunk(content=f"[Speaker notes] {notes}", page_number=slide_index + 1)
                )

    if slides_needing_render:
        rendered = _render_slide_images(file_path, len(prs.slides))
        for slide_index in slides_needing_render:
            image_bytes = rendered.get(slide_index)
            if image_bytes:
                images.append(
                    ExtractedImage(
                        image_id=str(uuid.uuid4()),
                        image_bytes=image_bytes,
                        page_number=slide_index + 1,
                        bbox=None,
                    )
                )

    return text_chunks, images
