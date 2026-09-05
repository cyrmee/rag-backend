import uuid

import docx

from app.extractors.types import ExtractedImage, ExtractionResult, TextChunk


def _table_to_markdown(table: docx.table.Table) -> str:
    rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
    if not rows:
        return ""

    lines = ["| " + " | ".join(rows[0]) + " |"]
    lines.append("| " + " | ".join("---" for _ in rows[0]) + " |")
    for row in rows[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)


def extract_docx(file_path: str) -> ExtractionResult:
    document = docx.Document(file_path)
    text_chunks: list[TextChunk] = []

    # docx has no true page concept at the XML level, so paragraph/table
    # index is used as a synthetic "page_number" throughout.
    for index, para in enumerate(document.paragraphs):
        if para.text.strip():
            text_chunks.append(TextChunk(content=para.text, page_number=index))

    for index, table in enumerate(document.tables):
        markdown = _table_to_markdown(table)
        if markdown:
            text_chunks.append(TextChunk(content=markdown, page_number=index))

    images: list[ExtractedImage] = []
    for index, rel in enumerate(document.part.related_parts.values()):
        if "image" in rel.content_type:
            images.append(
                ExtractedImage(
                    image_id=str(uuid.uuid4()),
                    image_bytes=rel.blob,
                    page_number=index,  # synthetic: relationship order, not a real page
                    bbox=None,
                )
            )

    return text_chunks, images
